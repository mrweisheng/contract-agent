# -*- coding: utf-8 -*-
"""API 路由。"""
import base64
import json
import os
import re
import tempfile
import traceback

from fastapi import APIRouter, BackgroundTasks, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from typing import Optional, List

from engine.types_config import TYPES, CURRENCY_OPTIONS, PORTS, get_type, template_path
from engine import fx
from engine.builder import build_contract, derive_payment
from engine.checker import check_rules, check_fingerprint, extract_contract_text
from engine import writer as W
from llm import client as llm
from llm.prompts import (extract_messages, parse_payment_messages, review_messages,
                         build_summary, settlement_struct_messages)
from api import store
from settlement import service as settlement_svc
from settlement import compute as settlement_compute

router = APIRouter(prefix="/api")

FILE_BIZ = {  # 生成文件名「日期+客户名+业务」里的业务段
    "car": "卖车", "transfer_gx": "高新过户", "transfer_ns": "纳税过户", "immigration": "高才移民",
}


def _gen_filename(req_type: str, form: dict) -> str:
    """生成文件名：日期+客户名+业务（如 20261010张三新办莲塘），比纯编号可读。
    编号仍单独分配并贯穿历史库/下载接口，文件名只影响落盘与展示。
    同客户同日重生成会撞名：已存在时加序号后缀（-2/-3），避免防覆盖归档
    把旧文件改名导致旧编号记录的下载 404。"""
    import datetime
    ymd = datetime.date.today().strftime("%Y%m%d")
    client = re.sub(r'[\\/:*?"<>|\s]+', "", str(form.get("client_name") or "")) or "客户"
    biz = FILE_BIZ.get(req_type, req_type)
    if req_type == "new_port":
        biz = "新办" + str(form.get("port") or "").replace("口岸", "").replace("大桥", "")
    base = f"{ymd}{client}{biz}"
    path = os.path.join(store.OUT_DIR, f"{base}.docx")
    n = 2
    while os.path.exists(path):
        path = os.path.join(store.OUT_DIR, f"{base}-{n}.docx")
        n += 1
    return os.path.basename(path)

# LLM 复核确认项过滤：结论为「检查通过」的条目不作为警告展示（提示词已禁，此处兜底）
NO_ISSUE_RE = re.compile(r"未发现.{0,8}(问题|异常|重复|缺失)")
NEG_MARK_RE = re.compile(r"不一致|未|缺|存在|但|错误|异常|应|建议|需|问题")


def _is_confirmation(iss: str) -> bool:
    if NO_ISSUE_RE.search(iss):
        return True
    return bool(re.search(r"一致|正确|无误", iss)) and not NEG_MARK_RE.search(iss)


# 解析 notes 噪音过滤：系统已自动处理的过程说明不算问题，不进提示/弹窗。
# 噪音＝「与预设不符/不一致/不同」（自定义分期是设计内路径）＋「…一致」类正常确认（"不一致"除外）
# ＋质保参数不全类提醒（客户没给时长/公里数时系统自动按默认「半年或3万公里」写入，无需确认）
# ＋「值已给却追问类型/类别」类废话（给了证件号还问是不是身份证——有值即可填表，无需客户确认）。
NOTE_NOISE_RE = re.compile(
    r"与.{0,4}预设.{0,12}(不符|不一致|不同)|(?<!不)一致"
    r"|质保.{0,30}(未说明|未提供|未明确|未写明|未给出|没说明|没有说明)"
    r"|未说明.{0,10}(类型|类别)|(类型|类别)未(说明|明确|写明|告知|提供)"
)


def _filter_notes(notes: list) -> list:
    return [n for n in (notes or []) if not NOTE_NOISE_RE.search(str(n))]


@router.get("/fx")
def get_fx():
    """前端费用输入行的实时约合提示：与合同落盘同一汇率源、同一加成口径。"""
    try:
        info = fx.get_hkd_cny()
    except fx.FXError as ex:
        raise HTTPException(502, str(ex))
    return {"rate": info["rate"], "source": info["source"]}


@router.get("/types")
def get_types():
    out = []
    for key, cfg in TYPES.items():
        out.append({
            "key": key,
            "label": cfg["label"],
            "short": cfg["short"],
            "currency_default": cfg["currency_default"],
            "currency_locked": bool(cfg.get("currency_locked")),
            "port_required": bool(cfg.get("port_required")),
            "ports": PORTS if cfg.get("port_required") else None,
            "groups": cfg["groups"],
            "fee_fields": cfg["fee_fields"],
            "pay_preset": cfg.get("pay_preset", []),
            "field_anchors": cfg.get("field_anchors"),  # 前端「聚焦字段→预览联动定位」锚点
            "ui_mode": cfg.get("ui_mode", ""),          # ledger＝事实台账布局（试点：高新现牌）
            "ledger_rows": cfg.get("ledger_rows"),      # 台账行组织（ui_mode=ledger 时用）
            "client_side": cfg["client_side"],
            "car_default": bool(cfg.get("car_default")),
            "imm_fee": bool(cfg.get("installment_table")),  # 移民式分期散文表：行号=期号（无总费用行）
        })
    return {"types": out, "currencies": CURRENCY_OPTIONS}


class ExtractReq(BaseModel):
    type: str
    text: str = Field(min_length=5)


@router.post("/extract")
def extract(req: ExtractReq):
    try:
        data = llm.chat_json(extract_messages(req.type, req.text.strip()))
    except llm.LLMError as ex:
        raise HTTPException(502, f"LLM 抽取失败：{ex}")
    fields = data.get("fields") or {}
    # 口岸纠正到合法值
    for k in ("port",):
        v = fields.get(k)
        if v and v not in PORTS:
            for p in PORTS:
                if v in p or p in v:
                    fields[k] = p
                    break
    return {
        "fields": fields,
        "currency": data.get("currency"),
        "total": data.get("total"),
        "payment_text": data.get("payment_text"),
        "notes": _filter_notes(data.get("notes")),
    }


class ParsePayReq(BaseModel):
    type: str
    currency: str
    total: int
    text: str = Field(min_length=3)
    sign_date: Optional[str] = None  # 签署日期,用于推算相对日期为绝对日期


@router.post("/parse-payment")
def parse_payment(req: ParsePayReq):
    try:
        data = llm.chat_json(parse_payment_messages(
            req.type, req.currency, req.total, req.text.strip(), req.sign_date or ""
        ))
    except llm.LLMError as ex:
        raise HTTPException(502, f"LLM 解析失败：{ex}")
    matches = bool(data.get("matches_preset"))
    result = {"matches_preset": matches, "notes": _filter_notes(data.get("notes"))}

    inst, extra_notes = [], []
    for x in (data.get("installments") or []):
        try:
            amt = int(str(x.get("amount", 0)).replace(",", ""))
        except ValueError:
            continue
        if amt <= 0:
            continue
        td = str(x.get("trigger_date") or "").strip() or None
        if td and not W.is_valid_iso_date(td):
            extra_notes.append(f"第{len(inst) + 1}期推算日期「{td}」不是有效日期，已忽略，请核对付款条件")
            td = None
        if td and req.sign_date and td < req.sign_date:
            extra_notes.append(f"第{len(inst) + 1}期付款日期 {td} 早于签署日期 {req.sign_date}，请核对")
        trigger = str(x.get("trigger") or "").strip()
        if not td and not trigger:
            extra_notes.append(f"第{len(inst) + 1}期未写明付款时间／节点／条件，请补充后再生成")
        inst.append({
            "seq": len(inst) + 1,
            "amount": amt,
            "label": str(x.get("label") or "").strip()[:6],
            "trigger": trigger,
            "trigger_date": td,
        })
    if extra_notes:
        result["notes"] = (result.get("notes") or []) + extra_notes
    preset_n = len(get_type(req.type).get("pay_preset") or [])
    if matches and len(inst) != preset_n:
        matches = False
        result["matches_preset"] = False  # 期数不符自动转自定义，属正常路径无需提示
    if not inst:
        raise HTTPException(422, "未解析出任何分期，请在付款计划表中手动填写")
    result["installments"] = inst
    result["sum"] = sum(x["amount"] for x in inst)
    return result


class PayRow(BaseModel):
    label: str = ""
    amount: int
    date: Optional[str] = None  # 绝对付款日期（YYYY-MM-DD）
    event: str = ""              # 付款事件条件


class Payment(BaseModel):
    installments: List[PayRow]
    source_text: Optional[str] = None


class GenerateReq(BaseModel):
    type: str
    form: dict
    payment: Payment


def _review_context(doc, cfg) -> str:
    """复核输入节选：双方信息表 + 费用表（含移民分期散文表）+ 付款条款区，
    舍去模板固定段落（指纹已逐一比对，无需占用复核注意力）。"""
    parts = []
    t = W.find_table(doc, cfg["party_table_anchor"])
    if t is not None:
        parts.append("【合约双方】\n" + "\n".join(
            " | ".join(c.text.strip() for c in row.cells) for row in t.rows))
    t = W.find_table(doc, cfg["fee_table_anchor"])
    if t is not None:
        parts.append("【服务费用】\n" + "\n".join(
            " | ".join(c.text.strip() for c in row.cells) for row in t.rows))
    if isinstance(cfg.get("installment_table"), str):
        # 落盘件用落盘态锚点（分期行改写后模板态锚点会失配）
        t = W.find_table(doc, cfg.get("installment_table_out") or cfg["installment_table"])
        if t is not None:
            parts.append("【付款计划】\n" + "\n".join(
                " | ".join(c.text.strip() for c in row.cells) for row in t.rows))
    try:
        s, e = W.section_range(doc, *cfg["payment_section"])
        body = "\n".join(p.text.strip() for p in doc.paragraphs[s:e] if p.text.strip())
        if body:
            parts.append("【付款条款】\n" + body)
    except ValueError:
        pass
    return "\n\n".join(parts)


def _run_review(no: str, type_key: str, form: dict, payment: dict, out_path: str):
    """后台 LLM 语义复核（软校验）：不阻塞生成主流程，结论写库供历史/轮询查询。
    thinking=False + 节选输入：MiniMax-M3 思考与答案共享 max_tokens，思考长度
    随机且耗尽即空输出重试，是原先耗时不可控的主因；复核属对照清单任务，
    无思考模式的结论质量足够（原末次重试兜底即为无思考）。"""
    try:
        from docx import Document
        cfg = get_type(type_key)
        doc = Document(out_path)
        text = _review_context(doc, cfg)
        rv = llm.chat_json(review_messages(text, build_summary(type_key, form, payment)),
                           max_tokens=1200, timeout=45, thinking=False)
        issues = [f"LLM复核：{i}" for i in (rv.get("issues") or []) if not _is_confirmation(i)]
        store.set_review(no, "pass" if not issues else "warn",
                         json.dumps(issues, ensure_ascii=False))
    except llm.LLMError as ex:
        store.set_review(no, "skip", json.dumps([f"LLM复核未完成（不影响合同使用）：{ex}"],
                                                ensure_ascii=False))
    except Exception as ex:
        traceback.print_exc()
        store.set_review(no, "skip", json.dumps([f"复核异常（不影响合同使用）：{ex}"],
                                                ensure_ascii=False))


@router.get("/review/{no}")
def get_review(no: str):
    """前端轮询后台复核结论。review_status：''=排队中 pass/warn=有结论 skip=未复核。"""
    rec = store.get_gen(no)
    if not rec:
        raise HTTPException(404, "记录不存在")
    try:
        issues = json.loads(rec.get("review_issues") or "[]")
    except Exception:
        issues = []
    return {"no": no, "review_status": rec.get("review_status") or "", "issues": issues}


@router.post("/generate")
def generate(req: GenerateReq, background_tasks: BackgroundTasks):
    cfg = get_type(req.type)
    form = req.form
    # 统一分期表 → 生成器付款 dict（服务端判定 模板预设/一次性/自定义，并校验分期合计=总费用）
    total = int(form.get("total_price") or form.get("total_fee") or 0)
    try:
        payment = derive_payment(req.type, req.payment.model_dump(), total)
    except ValueError as ex:
        raise HTTPException(422, str(ex))

    # 编号一律服务端自动分配（预约制：分配即落库占位，失败自动释放，无并发重复窗口）
    no = store.reserve_no()
    filename = _gen_filename(req.type, form)
    out_path = os.path.join(store.OUT_DIR, filename)

    try:
        report = build_contract(req.type, form, payment, no, out_path)
        report["agreement_no"] = no

        from docx import Document
        doc = Document(out_path)
        errors = check_rules(req.type, form, payment, report, doc)
        fp_issues = check_fingerprint(report["template"], out_path, req.type, report, form)
        errors.extend(fp_issues)

        warnings = []
        text = extract_contract_text(doc)
    except fx.FXError as ex:
        store.release_no(no)
        raise HTTPException(502, str(ex))
    except ValueError as ex:
        store.release_no(no)
        raise HTTPException(422, str(ex))
    except Exception as ex:
        traceback.print_exc()
        store.release_no(no)
        raise HTTPException(500, f"生成失败：{ex}")

    status = "ok" if not errors else "failed"
    store.confirm_no(no, req.type, form.get("client_name") or "", payment.get("mode", ""),
                     form.get("currency", ""), filename, status)
    # LLM 语义复核改后台异步：不阻塞生成主流程（原先同步等待 20~60 秒、最坏数分钟），
    # 结论写 gens 表供历史与 /api/review/{no} 轮询查询；failed 单已被拦截下载，无需复核
    if status == "ok":
        background_tasks.add_task(_run_review, no, req.type, form, payment, out_path)

    return {
        "ok": not errors,
        "no": no,
        "filename": filename,
        "errors": errors,
        "warnings": warnings,
        "text": text,
        "download_url": f"/api/download/{no}" if not errors else None,
    }


class PreviewReq(BaseModel):
    type: str
    form: dict = {}
    payment: dict = {}  # {installments: [{label, amount, date, event}], source_text}，录入中的宽松数据


@router.post("/preview")
def preview(req: PreviewReq):
    """实时 1:1 预览：与正式生成走同一个 builder（draft=True），返回真实 docx 字节。
    所见即所得是架构性保证——预览与下载产物是同一渲染管线的输出。
    隔离：不分配合同编号、不写 OUT_DIR、不落后端历史库、不跑 checker 与 LLM 复核。
    录入不完整不报错：占位渲染并在 draft_issues 里列出；仅模板结构损坏才 422。"""
    if req.type not in TYPES:
        raise HTTPException(422, f"未知合同类型 {req.type}")
    issues = []
    try:
        payment = derive_payment(req.type, req.payment or {}, None, draft=True, issues=issues)
        with tempfile.TemporaryDirectory(prefix="contract-preview-") as tmp:
            path = os.path.join(tmp, "preview.docx")
            report = build_contract(req.type, req.form or {}, payment, "（预览稿）", path,
                                    draft=True, issues=issues)
            with open(path, "rb") as f:
                data = f.read()
    except fx.FXError as ex:
        raise HTTPException(502, str(ex))
    except ValueError as ex:
        raise HTTPException(422, f"模板结构异常，无法预览：{ex}")
    except HTTPException:
        raise
    except Exception as ex:
        traceback.print_exc()
        raise HTTPException(500, f"预览渲染失败：{ex}")
    return {
        "docx_b64": base64.b64encode(data).decode(),
        "mode": report.get("mode"),
        "draft_issues": report.get("draft_issues") or [],
    }


@router.get("/download/{no}")
def download(no: str):
    rec = store.get_gen(no)
    if not rec:
        raise HTTPException(404, "记录不存在")
    if rec["status"] != "ok":
        raise HTTPException(403, "该合同未通过核对，不提供下载；请按问题清单修正后重新生成")
    path = os.path.join(store.OUT_DIR, rec["filename"])
    if not os.path.exists(path):
        raise HTTPException(404, "文件不存在")
    # 中文文件名双格式：filename* (RFC 5987) 供现代浏览器，ASCII 回退 filename 供
    # 部分 macOS/iOS Safari 与下载工具（Starlette 只发 filename*，回退缺失时名称乱码或下载失败）
    from urllib.parse import quote
    fallback = f"{no}{os.path.splitext(rec['filename'])[1]}"
    disposition = (f'attachment; filename="{fallback}"; '
                   f"filename*=utf-8''{quote(rec['filename'])}")
    return FileResponse(path, filename=None,
                        headers={"content-disposition": disposition},
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@router.get("/history")
def get_history():
    return {"items": store.history()}


# ==================== 收款结算单（独立轻流程，不接合同/编号） ====================

class SettlementParseReq(BaseModel):
    text: str = Field(min_length=5)


def _section_diff_notes(rule_data: dict, llm_data: dict) -> list:
    """规则解析 vs AI 整理的交叉校验：段落行数/金额对不上就点名段落请人工核对。
    规则在这里不当事实来源，只当抓 AI 漏行/幻行的网。
    规则一行都没读出的段落（全口语格式，规则读不了）不对比——AI 是唯一读者，没有 second opinion。"""
    labels = {"contract_total": "合同内总费用", "contract_paid": "合同内已收",
              "contract_unpaid": "合同内未收", "extra_fees": "合同外杂费",
              "extra_paid": "合同外杂费实收"}
    notes = []
    for key, label in labels.items():
        rl = rule_data.get(key) or []
        al = llm_data.get(key) or []
        if not rl:   # 规则没读出该段：无对比基准，不打扰
            continue
        if len(rl) == len(al) and abs(
                sum(float(r.get("amount") or 0) for r in rl)
                - sum(float(r.get("amount") or 0) for r in al)) <= 0.01:
            continue
        detail = f"（规则识别 {len(rl)} 行 / AI 识别 {len(al)} 行）" if len(rl) != len(al) else \
                 (f"（规则合计 {sum(float(r.get('amount') or 0) for r in rl):,.2f} / "
                  f"AI 合计 {sum(float(r.get('amount') or 0) for r in al):,.2f}）")
        notes.append(f"「{label}」段落规则解析与 AI 结果有差异{detail}，请重点核对该段")
    return notes


@router.post("/settlement/parse")
def settlement_parse(req: SettlementParseReq):
    """自由文本 → 业务结构化 JSON。
    主路：LLM 整理（对任意措辞/口语/标记稳定，禁止算术，只搬运数字与币种）；
    交叉校验：规则解析并行跑，段落行数/金额与 AI 不一致时点名请人工核对；
    抬头字段以规则解析逐字为准（防模型改写简繁用字）；
    兜底：LLM 不可用时回落规则解析并明示。算术全部留给确定性引擎。"""
    text = req.text.strip()
    rule_data = settlement_svc.parse_text(text)

    llm_data, llm_err = None, None
    try:
        raw = llm.chat_json(settlement_struct_messages(text), timeout=45)
        llm_data = settlement_svc.sanitize(raw)
        if not any(llm_data.get(k) for k in settlement_svc.SECTIONS):
            llm_data = None
            llm_err = "AI 没能从文本里整理出费用明细"
    except llm.LLMError as ex:
        llm_err = str(ex)

    if llm_data:
        data, source = llm_data, "llm"
        # 结构纠偏：规则解析当结构基准（段落归属、gift 标记），AI 当语义基准
        moved = settlement_svc.reconcile_structure(rule_data, data)
        # 抬头以规则解析为准：它逐字照抄原文，AI 会偶发把群名/客户改成繁体或改写。
        # 引擎不做任何繁简转换（输入什么就印什么），原文用字必须原样保留，不接受
        # 模型代转；规则没读出的字段（口语串联无「群名：」标签）才用 AI 的。
        for k in ("group", "customer", "date", "port", "settle_type", "date_iso"):
            if rule_data.get(k):
                data[k] = rule_data[k]
        notes = settlement_svc.parse_notes(data) + _section_diff_notes(rule_data, data)
        if moved:
            notes.insert(0, "「" + "、".join(moved) +
                         "」是抵扣/实收款项，已自动归入「合同外杂费实收」段，请核对")
    elif any(rule_data.get(k) for k in settlement_svc.SECTIONS):
        data, source = rule_data, "rules"
        notes = settlement_svc.parse_notes(data) + \
            [f"AI 整理暂不可用（{llm_err}），已按固定格式规则解析，请逐行核对"]
    else:
        raise HTTPException(
            422, f"没能识别结算单段落（{llm_err or '规则解析为空'}）。"
                 "请按「群名：…／合同内总费用。总价：…」的标准格式粘贴")
    return {"data": data, "notes": notes, "source": source}


@router.get("/settlement/rate")
def settlement_rate():
    """当天 CNY→HKD（5 位小数）。源与合同业务同一汇率提供器（新浪主源/腾讯备源）。"""
    try:
        return settlement_svc.today_rate()
    except fx.FXError as ex:
        raise HTTPException(502, str(ex))


class SettlementGenerateReq(BaseModel):
    data: dict
    rate: Optional[float] = None   # 用户手填 CNY→HKD，优先于系统汇率


@router.post("/settlement/preview")
def settlement_preview(req: SettlementGenerateReq):
    """模板稿预览：金额留空的行照列项目、金额留白；固定文件名覆盖，不存底。
    与正式出单同一渲染管线（与合同侧「预览=下载 1:1 同源」同一哲学）。"""
    try:
        return settlement_svc.preview(req.data, req.rate)
    except fx.FXError as ex:
        raise HTTPException(502, str(ex))
    except ValueError as ex:
        raise HTTPException(422, str(ex))
    except Exception as ex:
        traceback.print_exc()
        raise HTTPException(500, f"预览渲染失败：{ex}")


@router.post("/settlement/generate")
def settlement_generate(req: SettlementGenerateReq):
    data = req.data
    for field, label in (("group", "群名"), ("customer", "客户")):
        if not str(data.get(field) or "").strip():
            raise HTTPException(422, f"「{label}」未填，无法出单")
    if not any(data.get(k) for k in settlement_svc.SECTIONS):
        raise HTTPException(422, "没有任何费用明细，无法出单")
    try:
        return settlement_svc.generate(data, req.rate)
    except fx.FXError as ex:
        raise HTTPException(502, str(ex))
    except ValueError as ex:
        raise HTTPException(422, str(ex))
    except Exception as ex:
        traceback.print_exc()
        raise HTTPException(500, f"结算单生成失败：{ex}")


@router.get("/settlement/file/{kind}/{name}")
def settlement_file(kind: str, name: str):
    """成品文件：pdf/png。inline 输出（页面 <img>/<iframe> 直接用，下载靠 <a download>）。"""
    base = {"pdf": settlement_svc.PDF_DIR, "png": settlement_svc.PNG_DIR}.get(kind)
    if not base:
        raise HTTPException(404, "未知文件类型")
    safe = os.path.basename(name)
    if safe != name or not safe:
        raise HTTPException(404, "文件不存在")
    path = os.path.join(base, safe)
    if not os.path.isfile(path):
        raise HTTPException(404, "文件不存在")
    from urllib.parse import quote
    media = "application/pdf" if kind == "pdf" else "image/png"
    disposition = (f'inline; filename="{safe.encode("ascii", "ignore").decode() or "file"}"; '
                   f"filename*=utf-8''{quote(safe)}")
    return FileResponse(path, headers={"content-disposition": disposition}, media_type=media)


@router.get("/health")
def health():
    return {"status": "ok", "model": llm.MODEL, "key_set": bool(llm.API_KEY)}
