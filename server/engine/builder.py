# -*- coding: utf-8 -*-
"""
合同组装器：按类型配置 + 表单数据 + 付款计划，在模板副本上生成合同。
确定性渲染（条款文字由模板句式生成，LLM 不参与落盘文字），保证可控。
"""
import copy
import datetime
import os
import re
import shutil
from docx import Document

from engine import writer as W
from engine.types_config import get_type, template_path, TEMPLATES_DIR
from engine import fx
from engine.money import (to_cn_upper, with_commas, currency_label, currency_symbol,
                          amount_phrase, amount_phrase_dual, dual_numeric,
                          clean_amount, clean_currency)

PARTY_LABEL_MAP = [
    ("证件", "client_id"),
    ("联络人", "client_contact"),
    ("联系人", "client_contact"),
    ("联络电话", "client_phone"),
    ("电话", "client_phone"),
    ("联络地址", "client_address"),
    ("地址", "client_address"),
]


def _party_label_to_field(label_text: str):
    if "买方" in label_text or "委托方" in label_text or "甲方（委托方" in label_text:
        return "client_name"
    for kw, field in PARTY_LABEL_MAP:
        if kw in label_text:
            return field
    return None


PLACEHOLDER_DATE = "（日期待填）"


def _lenient_amount(value, name, draft, issues, default=0):
    """非 draft 与 clean_amount 行为完全一致；draft 下解析失败记入 issues 并用默认值，
    负数钳为 0——预览绝不因录入问题中断，真生成仍严格拦截。"""
    try:
        n = clean_amount(value, name)
    except ValueError as ex:
        if not draft:
            raise
        issues.append(str(ex))
        return default
    if draft and n < 0:
        issues.append(f"{name}不能为负，预览按 0 显示")
        return 0
    return n


def _date_cn_safe(iso, label, draft, issues):
    """有效日期 → date_cn；draft 下无效/空缺记入 issues 并返回占位文字；
    非 draft 无效值严格拦截（与真生成同口径的 ValueError，而非下游崩溃）。"""
    s = str(iso or "").strip()
    if W.is_valid_iso_date(s):
        return W.date_cn(s)
    if not draft:
        raise ValueError(f"{label}「{s}」不是有效日期（YYYY-MM-DD）")
    issues.append(f"{label}「{s}」无法识别" if s else f"{label}待填")
    return PLACEHOLDER_DATE


def _archive_existing(out_path: str) -> str:
    """目标文件已存在时，先归档旧件（改名加时间戳）再写入，绝不静默覆盖。"""
    base, ext = os.path.splitext(out_path)
    stamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    archived = f"{base}-归档{stamp}{ext}"
    n = 1
    while os.path.exists(archived):
        archived = f"{base}-归档{stamp}-{n}{ext}"
        n += 1
    shutil.move(out_path, archived)
    return archived


# 签署块日期空位：年份可预填（高才移民模板「日期：2026年____月____日」），填入时整段替换
SIG_DATE_BLANK_RE = re.compile(r"日期：(?:_{2,}|\d{4})年_{2,}月_{2,}日")
# 「日期：」标签格与空位格分离的模板（高才移民）：空位独占一段，无「日期：」前缀
SIG_DATE_BARE_RE = re.compile(r"(?:_{2,}|\d{4})年_{2,}月_{2,}日")


def _fill_signature_dates(doc, y, m, d) -> int:
    """签署确认块双方日期空位精确替换（只动日期空位，签名/姓名/印章留给手签）。
    两种模板形态：「日期：____年__月__日」同段；「日期：」标签+「2026年__月__日」独立格。
    不做已填去重——填过的段落不再匹配空位模式，天然幂等（曾用 id(p._p) 去重，
    但 lxml 代理对象 GC 后 id 会复用，会把另一处同文空位误判成已填导致漏填）。"""
    filled = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    mm = SIG_DATE_BLANK_RE.search(p.text)
                    if mm:
                        old, new = mm.group(0), f"日期：{y}年{m}月{d}日"
                    else:
                        bm = SIG_DATE_BARE_RE.fullmatch(p.text.strip())
                        if not bm:
                            continue
                        old, new = bm.group(0), f"{y}年{m}月{d}日"
                    if W.replace_in_paragraph(p, old, new):
                        filled += 1
    return filled


def build_contract(type_key: str, form: dict, payment: dict, agreement_no: str, out_path: str,
                   draft: bool = False, issues: list = None, fx_rate: float = None):
    """生成合同，返回变更报告（供核对器比对）。
    draft=True（实时预览）：录入类问题（金额/日期/币种/质保参数）不中断，
    用占位值渲染并汇总到 report["draft_issues"]；模板结构类问题（锚点/表格/
    空位缺失）仍然抛 ValueError——结构坏了预览也无意义。"""
    if draft and issues is None:
        issues = []
    cfg = get_type(type_key)
    try:
        tpl = template_path(cfg, form)
    except ValueError as ex:
        # 录入类问题（口岸未选）：预览回退到该类型第一个模板，保证有画面可看
        if not draft:
            raise
        issues.append(str(ex))
        tpl = os.path.join(TEMPLATES_DIR, next(iter(cfg["template_map"].values())))
    # 入参校验前置：币种与总费用先解析，任一不合法就在复制模板之前失败，
    # 避免在 out/ 留下不在历史列表里的孤儿模板副本。
    raw_cur = (cfg["currency_default"] if cfg.get("currency_locked")
               else (form.get("currency") or cfg["currency_default"]))
    try:
        cur = clean_currency(raw_cur)
    except ValueError as ex:
        if not draft:
            raise
        cur = cfg["currency_default"]
        issues.append(f"币种不支持，预览按默认币种渲染：{ex}")
    total = _lenient_amount(form.get("total_price") or form.get("total_fee") or 0,
                            "总费用", draft, issues)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    # 纵深防御：同号文件若已存在，先归档旧件，避免旧合同被静默覆盖
    if os.path.exists(out_path):
        _archive_existing(out_path)
    shutil.copyfile(tpl, out_path)
    doc = Document(out_path)

    report = {"type": type_key, "currency": cur, "fee_rows": [], "section_paras": None}
    if fx_rate is None:
        fx_info = fx.get_hkd_cny()
        rate = fx_info["rate"]
        report["fx_rate"], report["fx_source"] = rate, fx_info["source"]
    else:
        rate = fx_rate
        report["fx_rate"], report["fx_source"] = rate, "fixed"

    # 1. 顶部编号/日期行（meta_number_line：编号独立成段且无日期空位，只填编号；
    #    签署日期只落在签署确认块，与 _fill_signature_dates 同源）
    sign_date = str(form.get("sign_date") or "").strip()
    if sign_date and draft and not W.is_valid_iso_date(sign_date):
        issues.append(f"签署日期「{sign_date}」无法识别，预览按今日显示")
        sign_cn = W.today_cn()
    else:
        sign_cn = W.date_cn(sign_date) if sign_date else W.today_cn()
    if cfg.get("meta_number_line"):
        meta = W.find_paragraph(doc, cfg["meta_number_line"])
        if meta is None or W.fill_blanks_in_paragraph(meta, [agreement_no]) < 1:
            raise ValueError("模板编号行缺少空位，请检查模板")
    else:
        meta = doc.paragraphs[0]
        n = W.fill_blanks_in_paragraph(meta, [agreement_no, sign_cn])
        if n < 2:
            raise ValueError("模板顶部编号/日期行缺少空位，请检查模板")

    # 2. 甲乙双方表（客户侧）
    party_tbl = W.find_table(doc, cfg["party_table_anchor"])
    if party_tbl is None:
        raise ValueError("找不到甲乙双方信息表")
    lc, vc = cfg["party_label_col"], cfg["party_value_col"]
    for row in party_tbl.rows:
        label = row.cells[lc].text.strip()
        field = _party_label_to_field(label)
        if field:
            val = str(form.get(field) or "").strip()
            W.set_cell_text(row.cells[vc], val)

    # 3. 键值表（车辆/指标数据）
    for kv in cfg.get("kv_tables", []):
        t = W.find_table(doc, kv["anchor"])
        if t is None:
            raise ValueError(f"找不到数据表（锚点 {kv['anchor']}）")
        for row in t.rows:
            label = row.cells[0].text.strip()
            for key, field in kv["map"].items():
                if key in label:
                    val = form.get(field)
                    if field == "quota_type":
                        val = cfg["quota_auto"]
                    W.set_cell_text(row.cells[1], str(val or "").strip() or "—")

    # 3.5 申请人信息表（高才移民）：整行重写为模板同款句式；
    #     空值落全角空位（打印后可手写），主申请人缺省同委托方
    app_cfg = cfg.get("applicant_table")
    if app_cfg:
        t = W.find_table(doc, app_cfg["anchor"])
        if t is None:
            raise ValueError("找不到申请人信息表")
        lead = (str(form.get("lead_applicant") or "").strip()
                or str(form.get("client_name") or "").strip() or "　　")
        spouse = str(form.get("spouse_name") or "").strip() or "　　"
        minors = str(form.get("minor_children") or "").strip() or "　"
        totaln = str(form.get("total_applicants") or "").strip() or "　"
        if draft and not any(str(form.get(k) or "").strip() for k in
                             ("lead_applicant", "spouse_name", "minor_children", "total_applicants")):
            issues.append("申请人信息未填：主申请人按委托方、随行人数留空位（可打印后手写）")
        W.set_paragraph_text(t.rows[0].cells[0].paragraphs[0],
                             app_cfg["lead_line"].format(name=lead))
        W.set_paragraph_text(t.rows[1].cells[0].paragraphs[0],
                             app_cfg["family_line"].format(spouse=spouse, minors=minors, total=totaln))

    # 4. 换车费用（过户类）。行名默认「换车费用」，客户说「上车费用」则随客户用词
    #    （两名词同指现牌流程中车辆相关费用，白名单防幻觉，其余输入一律回默认）
    if cfg.get("exchange_fee"):
        t = W.find_table(doc, cfg["exchange_fee"]["anchor"])
        if t is not None:
            amt = _lenient_amount(form.get("exchange_fee") or 0, "换车费用", draft, issues)
            xlabel = str(form.get("exchange_label") or "").strip()
            if xlabel not in (cfg["exchange_fee"]["anchor"], *cfg["exchange_fee"].get("aliases", [])):
                xlabel = cfg["exchange_fee"]["anchor"]
            row = t.rows[0]
            W.set_cell_text(row.cells[0], xlabel)
            W.set_cell_text(row.cells[1], amount_phrase_dual(amt, cur, rate))
            report["exchange_fee"] = amt
            report["exchange_label"] = xlabel

    # 5. 付款计划 → 条款区 + 费用表
    fee_tbl = W.find_table(doc, cfg["fee_table_anchor"])
    if fee_tbl is None:
        raise ValueError("找不到费用表")
    report["total"] = total  # total 已在函数开头严格解析（P1-4：绝不静默截断）
    mode = payment.get("mode", "default")
    report["mode"] = mode
    sec_no, next_no = cfg["payment_section"]

    if cfg.get("installment_table"):
        _immigration_fee(doc, cfg, fee_tbl, payment, cur, total, report, draft, issues, rate, mode)
    elif mode == "default":
        _default_mode(doc, cfg, fee_tbl, form, payment, cur, total, report, sign_cn, draft, issues, rate)
    elif mode == "one_time":
        _one_time_mode(doc, cfg, fee_tbl, payment, cur, total, report, sec_no, next_no, draft, issues, rate)
    elif mode == "custom":
        _custom_mode(doc, cfg, fee_tbl, payment, cur, total, report, sec_no, next_no, draft, issues, rate)
    else:
        raise ValueError(f"未知付款模式 {mode}")

    # 6. 卖车可选内容：附赠项 / 发动机质保（客户信息提到才落盘，默认全无）
    if type_key == "car":
        _apply_car_extras(doc, form, report, draft, issues)

    # 7. 签署确认块：双方签署日期（与顶部签署日期同源：录入日默认今日，可手动改期）
    if W.is_valid_iso_date(sign_date):
        y, m, d = W.iso_date_parts(sign_date)
    else:
        t = datetime.date.today()
        y, m, d = str(t.year), str(t.month), str(t.day)
    filled = _fill_signature_dates(doc, y, m, d)
    if filled < 2:
        msg = "签署确认块日期空位不足（应甲乙双方各一处），请检查模板"
        if not draft:
            raise ValueError(msg)
        issues.append(msg)
    report["signature_dates"] = filled

    doc.save(out_path)
    report["template"] = tpl
    report["output"] = out_path
    if draft:
        report["draft_issues"] = issues
    return report


# ---------------- 统一付款表 → 生成器付款 dict ----------------

def _apply_car_extras(doc, form: dict, report: dict, draft: bool, issues: list):
    """合并版卖车模板的条件块（server/templates/1-卖车，中文条款号）：
    A. 四节尾「甲方陈述并保证」块——承诺无重大事故才保留；
    B. 五节附赠服务——有附赠行才保留并写入推导文字；
    C. 六节售后质保——开启质保才保留，自定义参数写入保修范围句；
    最后全量重排中文条款号。赠过户费/车险时最小改写八节对应句。文字推导统一走 car_extras。"""
    from engine import car_extras as X

    gifts = X.gift_lines(form)
    try:
        wp = X.warranty_paragraph(form)
    except ValueError as ex:
        if not draft:
            raise
        issues.append(str(ex))
        # 参数无法识别：预览按模板默认期限渲染，让用户至少看到质保节形态
        wp = (X.WARRANTY_TMPL.format(a=X.DEFAULT_MONTHS_PHRASE, b=X.DEFAULT_KM_PHRASE)
              if X.warranty_on(form) else None)

    # A. 陈述与保证块（「甲方陈述并保证：」起至下一标题）
    if not X.accident_free_on(form):
        i = W.heading_index(doc, "甲方陈述并保证：")
        j = W.next_heading_index(doc, i)
        W.delete_span(doc, i, j)

    # B. 附赠服务节
    i = W.heading_index(doc, "附赠服务")
    j = W.next_heading_index(doc, i)
    if gifts:
        W.rewrite_span(doc, i + 1, j, gifts)
    else:
        W.delete_span(doc, i, j)

    # C. 售后质保节
    i = W.heading_index(doc, "售后质保")
    j = W.next_heading_index(doc, i)
    if X.warranty_on(form):
        body = doc.paragraphs[i + 1:j]
        target = next((p for p in body if "保修范围" in p.text), None)
        if target is None:
            raise ValueError("质保节缺少保修范围句，请检查模板")
        try:
            period, km = X.warranty_custom_parts(form)
        except ValueError:
            period = km = None  # 参数非法：draft 已记 issue，按模板默认渲染
        if period:
            W.replace_in_paragraph(target, "六个月", period)
        if km:
            W.replace_in_paragraph(target, "三万公里", km)
    else:
        W.delete_span(doc, i, j)
    report["gift_lines"] = gifts
    report["warranty_on"] = bool(wp)
    report["accident_free"] = X.accident_free_on(form)

    # 八节：赠过户费/车险时改写过户费用句/保险句（逐字锚点，禁止盲改）
    if X.gift_transfer_on(form):
        p_pass = next((p for p in doc.paragraphs if "车辆过户手续由甲方负责办理" in p.text), None)
        if p_pass is None:
            raise ValueError("找不到过户费用句，请检查模板")
        if p_pass.text.strip() != X.PASSAGE_ORIG:
            raise ValueError("模板过户费用句原文与预设不符，禁止改写，请检查模板")
        W.set_paragraph_text(p_pass, X.passage_text(form))
    if X.gift_insurance(form):
        p_ins = next((p for p in doc.paragraphs if "现有保险不随车转移" in p.text), None)
        if p_ins is None:
            raise ValueError("找不到保险句，请检查模板")
        if p_ins.text.strip() != X.INS_ORIG:
            raise ValueError("模板保险句原文与预设不符，禁止改写，请检查模板")
        W.set_paragraph_text(p_ins, X.insurance_text(form))

    W.renumber_headings(doc)


def _node_match(x: dict, p: dict) -> bool:
    """分期行与模板预设行是否同一节点：事件行须与预设原文一致；签约当日期行有日期即可。"""
    ev = (x.get("trigger") or "").strip()
    pe = (p.get("event") or "").strip()
    if pe:
        return ev == pe
    if ev and "签约" not in ev:
        return False
    return W.is_valid_iso_date(x.get("trigger_date")) or bool(ev)


def derive_payment(type_key: str, payment: dict, total: int = None,
                   draft: bool = False, issues: list = None) -> dict:
    """前端统一分期表 → 生成器所需的付款 dict。
    判定：与模板预设逐期一致 → default（保持模板条款，仅填金额）；
    单期 → one_time；其余 → custom（重写条款为一句总括）。

    total：可选，传入总售价/总费用时，对「尾款 = 总额 − 订金」做一致性校验（P2-10 纵深防御）。
    draft：预览用。录入不完整不拦截，问题记入 issues，字段按 0/空兜底；
    非 draft 行为与报错文案保持原样。
    """
    if draft and issues is None:
        issues = []
    cfg = get_type(type_key)
    preset = cfg.get("pay_preset") or []
    inst = []
    for i, x in enumerate(payment.get("installments") or []):
        raw = x.get("amount", 0)
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            msg = f"第{i + 1}期金额为空，请填写后再生成"
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
            amt = 0
        else:
            try:
                amt = clean_amount(raw, f"第{i + 1}期金额")
            except ValueError as ex:
                if not draft:
                    raise ValueError(str(ex))
                issues.append(str(ex))
                amt = 0
        if amt <= 0:
            msg = f"第{i + 1}期金额必须大于 0"
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
            amt = max(amt, 0)
        inst.append({
            "seq": len(inst) + 1,
            "amount": amt,
            "label": (x.get("label") or "").strip()[:6],
            "trigger": (x.get("event") or "").strip(),
            "trigger_date": (x.get("date") or "").strip() or None,
        })
    if not inst:
        if not draft:
            raise ValueError("缺少付款计划（至少一期）")
        issues.append("缺少付款计划（至少一期）")
        return {"mode": "default", "source_text": payment.get("source_text")}
    # P2-10 纵深防御：各期合计必须等于总额（覆盖 default/one_time/custom 全部模式）。
    # 否则尾款被引擎按「总额 − 订金」重算、或自定义分期写错时，合同金额会与用户输入不符且无人知晓。
    if total is not None and not draft:
        s = sum(x["amount"] for x in inst)
        if s != total:
            raise ValueError(
                f"分期合计（{s}）与总费用（{total}）不一致，请核对各期金额")
    for i, x in enumerate(inst):
        if not (W.is_valid_iso_date(x["trigger_date"]) or x["trigger"]):
            msg = f"第{i + 1}期缺少付款日期／事件条件"
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
    src = payment.get("source_text")

    if preset and len(inst) == len(preset) and all(_node_match(x, p) for x, p in zip(inst, preset)):
        legacy = {"mode": "default", "source_text": src}
        if cfg.get("car_default"):
            d1 = inst[0]["trigger_date"]
            if not W.is_valid_iso_date(d1):
                if not draft:
                    raise ValueError("卖车订金须填写付款日期（签约当日即选签署日期）")
                issues.append("卖车订金须填写付款日期（签约当日即选签署日期）")
            else:
                legacy["deposit_date"] = d1
            legacy["deposit_amount"] = inst[0]["amount"]
            d2 = inst[1]["trigger_date"]
            if W.is_valid_iso_date(d2):
                legacy["balance_date"] = d2
            else:
                # 客户未给尾款日期：不造日期，保留事件改写条款
                legacy["balance_event"] = inst[1]["trigger"]
        else:
            legacy.update(pay1=inst[0]["amount"], pay2=inst[1]["amount"], pay3=inst[2]["amount"])
        return legacy
    if len(inst) == 1:
        return {"mode": "one_time", "pay_date": inst[0]["trigger_date"],
                "pay_event": inst[0]["trigger"], "source_text": src}
    return {"mode": "custom", "installments": inst, "source_text": src}


# ---------------- 三种付款模式 ----------------

# ---------------- 付款时点措辞（custom 分期条款 / 移民分期散文共用） ----------------

SIGN_DAY_RE = re.compile(r"^(本合约|本协议|合同)?签订(当日|当天|之日)$|^签约(当日|当天|之日|时)$")
PAY_END_RE = re.compile(r"(支付|付清|付款|支付完毕|付)$")
# 时点短语以「定金/尾款/款项」等宾语结尾时自带动词（支付定金/付尾款），不再补「支付」
PAY_OBJ_RE = re.compile(r"[金款]$")


def pay_when_phrase(when: str) -> str:
    """「于{when}支付…」句式用的时点短语：不以「付」类动词或「金/款」宾语结尾才补「支付」。"""
    if PAY_END_RE.search(when) or PAY_OBJ_RE.search(when):
        return when
    return when + "支付"


def payment_when(ins) -> str:
    """一期付款时点文字：绝对日期优先；「签约当日」类短语与日期同源只留日期；
    描述里已含同一日期直接用原话（避免「X日或X日…」）；其余返回条件原文。"""
    iso = None
    td = (ins.get("trigger_date") or "").strip()
    if W.is_valid_iso_date(td):
        iso = td
    else:
        # trigger_date 缺失/非法时，从条件描述里捞「X年X月X日」
        m = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", (ins.get("trigger") or "").strip())
        if m:
            cand = f"{int(m.group(1))}-{int(m.group(2))}-{int(m.group(3))}"
            if W.is_valid_iso_date(cand):
                iso = cand
    ev = (ins.get("trigger") or "").strip()
    if iso and ev:
        if W.date_cn(iso) in ev:
            return ev
        if not SIGN_DAY_RE.match(re.sub(r"(支付|付清|付款|支付完毕)$", "", ev)):
            return f"{W.date_cn(iso)}或{ev}"
        return W.date_cn(iso)
    return W.date_cn(iso) if iso else ev


def _rewrite_payment(doc, cfg, report, sec_no, next_no, texts):
    """改写条款区付款段（只动分期段，保留费用范围/逾期/收款账户等无关条款）。"""
    labels = [p["label"] for p in (cfg.get("pay_preset") or [])]
    total, a, b = W.rewrite_section(doc, sec_no, next_no, texts, labels or None)
    report["section_paras"] = total
    report["payment_rewrite"] = (a, b)


def _row_name(installments, i: int) -> str:
    """分期行名（与解析回填共用 writer.payment_row_name，所见即所得）。"""
    labels = [str(x.get("label") or "").strip() for x in installments]
    return W.payment_row_name(labels, i)


def _fill_amount_cell(row, cfg, cur, amount, rate):
    style = cfg["fee_cell_style"]
    if style == "split":
        W.set_cell_text(row.cells[1], f"{currency_label(cur)}（大写）{to_cn_upper(amount)}")
        W.set_cell_text(row.cells[2], dual_numeric(amount, cur, rate))
    else:  # single / single_rmb 统一替换单元格为完整金额短语（含第二币种换算）
        W.set_cell_text(row.cells[1], amount_phrase_dual(amount, cur, rate))


def _default_mode(doc, cfg, fee_tbl, form, payment, cur, total, report, sign_cn, draft, issues, rate):
    # 费用表：总金额行
    rows = fee_tbl.rows
    _fill_amount_cell(rows[0], cfg, cur, total, rate)

    if cfg.get("car_default"):
        # 卖车默认两段：订金 / 订金日期 / 尾款（自动=总价-订金）
        deposit = _lenient_amount(payment.get("deposit_amount") or 0, "订金金额", draft, issues)
        if not 0 < deposit < total:
            msg = (f"订金金额须大于 0 且小于总售价（当前订金 {deposit}，总价 {total}；"
                   f"一次性付清请用「一次性付款」模式）")
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
        balance = total - deposit
        if draft and balance <= 0:
            balance = 0
        deposit = max(deposit, 0)
        _fill_amount_cell(rows[1], cfg, cur, deposit, rate)
        W.set_cell_text(rows[2].cells[1],
                        _date_cn_safe(payment.get("deposit_date"), "订金日期", draft, issues))
        _fill_amount_cell(rows[3], cfg, cur, balance, rate)
        report["fee_rows"] = [W.row_label(r) for r in rows]
        report["amounts"] = [total, deposit, balance]
        # 三节尾款句：金额 + 日期空位顺序填空（客户给了尾款日期 → 填空保留「日期或过户当日」；
        # 只有事件 → 整句改写为单事件句含金额，绝不编造日期）
        start, end = W.section_range(doc, *cfg["payment_section"])
        p = doc.paragraphs[start]
        if "剩余购车尾款" not in p.text:
            raise ValueError("卖车三节尾款句与模板不符，请检查模板")
        bd = payment.get("balance_date")
        if W.is_valid_iso_date(str(bd or "").strip()):
            y, m, d = W.iso_date_parts(bd)
            cn = to_cn_upper(balance)
            if cn.endswith("元整"):   # 模板空位后自带「元整」，大写值只填金额主体
                cn = cn[:-2]
            filled = W.fill_blanks_in_paragraph(
                p, [cn, with_commas(balance), y, m, d])
            if filled < 5 or BLANK_LEFT(p):
                raise ValueError("卖车三节尾款句空位数量异常或填空后仍有残留")
            if cur != "HKD":   # 尾款句模板内置「港币/HK$」字面量，非港币合同须随币种改写
                W.replace_in_paragraph(p, "港币", currency_label(cur))
                W.replace_in_paragraph(p, "HK$", currency_symbol(cur))
        else:
            ev = (payment.get("balance_event") or "").strip() \
                or (cfg["pay_preset"][1].get("event") or "车辆完成香港运输署过户登记手续当日")
            W.set_paragraph_text(
                p, f"剩余购车尾款{currency_label(cur)}{to_cn_upper(balance)}"
                   f"（{currency_symbol(cur)}{with_commas(balance)}），须于{ev}，由乙方一次性支付予甲方。")
        report["section_paras"] = end - start
    else:
        # 过户/新办默认三期（触发时点为模板固定文字，仅填金额）
        amounts = [_lenient_amount(payment.get(k) or 0, f"第{i + 1}期金额", draft, issues)
                   for i, k in enumerate(("pay1", "pay2", "pay3"))]
        for i, amt in enumerate(amounts):
            _fill_amount_cell(rows[i + 1], cfg, cur, amt, rate)
        report["fee_rows"] = [W.row_label(r) for r in rows]
        report["amounts"] = [total] + amounts
        # 定金时点展开为绝对日期（生成日即签约日，与顶部签署日期同源）：
        # 新办「本协议签订后即日支付」/ 过户「于本合约签订当日支付」
        start, end = W.section_range(doc, *cfg["payment_section"])
        for p in doc.paragraphs[start:end]:
            for old, new in (
                ("本协议签订后即日支付", f"本协议签订当日（{sign_cn}）支付"),
                ("于本合约签订当日支付", f"于本合约签订当日（{sign_cn}）支付"),
            ):
                if old in p.text:
                    W.replace_in_paragraph(p, old, new)
                    break
        report["section_paras"] = None  # 条款文字不变（仅定金句内替换日期）


def BLANK_LEFT(p) -> bool:
    return bool(re.search(r"_{2,}", "".join(r.text for r in p.runs)))


def _ot_desc(payment, draft, issues) -> str:
    date = payment.get("pay_date") or ""
    event = (payment.get("pay_event") or "").strip()
    date_cn = _date_cn_safe(date, "付款日期", draft, issues) if date else ""
    if date_cn and event:
        return f"{date_cn}或{event}"
    if date_cn:
        return date_cn
    return event or ""


def _one_time_mode(doc, cfg, fee_tbl, payment, cur, total, report, sec_no, next_no, draft, issues, rate):
    rows = fee_tbl.rows
    _fill_amount_cell(rows[0], cfg, cur, total, rate)
    desc = _ot_desc(payment, draft, issues)
    if not desc:
        msg = "一次性付清需提供支付日期或条件"
        if not draft:
            raise ValueError(msg)
        issues.append(msg)
        desc = PLACEHOLDER_DATE
    W.set_cell_text(rows[1].cells[0], "付款方式")
    if cfg["fee_cell_style"] == "split":
        W.set_cell_text(rows[1].cells[1], f"一次性全款支付（{desc}）")
        W.set_cell_text(rows[1].cells[2], "")
    else:
        W.set_cell_text(rows[1].cells[1], f"一次性全款支付（{desc}）")
    # 删除其余分期行
    for idx in range(len(rows) - 1, 1, -1):
        W.delete_row(fee_tbl, idx)
    report["fee_rows"] = [W.row_label(r) for r in fee_tbl.rows]
    report["amounts"] = [total]
    # 条款改写为单段
    payer, payee = cfg["payer"], cfg["payee"]
    if cfg.get("car_default"):
        text = (f"车辆总售价{amount_phrase(total, cur)}须于{desc}，"
                f"由乙方一次性支付予甲方。")
    else:
        text = (f"服务总费用{amount_phrase(total, cur)}须于{desc}，"
                f"由{payer}一次性支付予{payee}。")
    _rewrite_payment(doc, cfg, report, sec_no, next_no, [text])


def _custom_mode(doc, cfg, fee_tbl, payment, cur, total, report, sec_no, next_no, draft, issues, rate):
    installments = payment.get("installments") or []
    if not installments:
        raise ValueError("自定义分期缺少期次数据")
    amounts = [_lenient_amount(x.get("amount") or 0, f"第{i + 1}期金额", draft, issues)
               for i, x in enumerate(installments)]
    if any(a <= 0 for a in amounts):
        msg = "各期金额必须大于 0"
        if not draft:
            raise ValueError(msg)
        issues.append(msg)
        amounts = [max(a, 0) for a in amounts]
    if sum(amounts) != total:
        msg = f"各期合计 {sum(amounts)} ≠ 总费用 {total}，请先在界面修正"
        if not draft:
            raise ValueError(msg)
        issues.append(msg)
    for i, ins in enumerate(installments):
        has_date = W.is_valid_iso_date((ins.get("trigger_date") or "").strip())
        has_cond = bool((ins.get("trigger") or "").strip())
        if not (has_date or has_cond):
            msg = f"第{i + 1}期缺少付款时间／节点／条件（既无日期也无触发条件），不能生成"
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
    rows = fee_tbl.rows
    _fill_amount_cell(rows[0], cfg, cur, total, rate)
    # 统一行样式：以模板第1数据行为源，删除原有分期行后克隆 n 份
    src_tr = copy.deepcopy(rows[1]._tr)
    while len(fee_tbl.rows) > 1:
        W.delete_row(fee_tbl, 1)
    for _ in range(len(installments)):
        fee_tbl._tbl.append(copy.deepcopy(src_tr))
    rows = fee_tbl.rows
    for i, ins in enumerate(installments):
        row = rows[i + 1]
        W.set_cell_text(row.cells[0], _row_name(installments, i))
        _fill_amount_cell(row, cfg, cur, amounts[i], rate)
    report["fee_rows"] = [W.row_label(r) for r in rows]
    report["amounts"] = [total] + amounts

    # 付款时间/事件逐期写入条款区文字（表格保持模板形态只列名称+金额，不加日期列），
    # 行名与费用表行名一致，一一对应；预设节点的专业文案由解析层回填进 trigger
    texts = []
    for i, ins in enumerate(installments):
        when = payment_when(ins) or PLACEHOLDER_DATE   # 占位仅 draft 可达，非 draft 已被 derive 拦截
        texts.append(f"{_row_name(installments, i)}：于{pay_when_phrase(when)}。")
    _rewrite_payment(doc, cfg, report, sec_no, next_no, texts)


# ---------------- 移民式费用区（总费用行表 + 分期散文表） ----------------

def _imm_set_row(row, label, verb, amt, cur, rate):
    """分期散文行整句改写（沿用模板散文句式）：主币大写 + 双币数字对，与两地牌
    amount_phrase_dual 同款，如「人民币（大写）壹拾万元整（¥100,000 / HK$117,716）」。"""
    W.set_paragraph_text(row.cells[1].paragraphs[0],
                         f"{label}  {verb}{amount_phrase_dual(amt, cur, rate)}。")


def _immigration_fee(doc, cfg, fee_tbl, payment, cur, total, report, draft, issues, rate, mode):
    """高才移民模板费用区，与其他业务同款双币逻辑：
    总费用行（模板 4 列，对调后 C1=主币大写、C2=双币数字）+ 分期散文表（三模式统一
    整句改写，行内含主币金额/大写/约合他币）+「分N期支付」格随期数改写。"""
    if fee_tbl is None:
        raise ValueError("找不到费用表")
    row0 = fee_tbl.rows[0]
    W.set_cell_text(row0.cells[1], f"{currency_label(cur)}（大写）{to_cn_upper(total)}")
    W.set_cell_text(row0.cells[2], dual_numeric(total, cur, rate))

    inst_tbl = W.find_table(doc, cfg["installment_table"])
    if inst_tbl is None:
        raise ValueError("找不到分期付款表")
    amounts = []

    if mode == "default":
        preset = cfg.get("pay_preset") or []
        if len(inst_tbl.rows) != len(preset):
            raise ValueError("分期付款表行数与预设不符，请检查模板")
        amounts = [_lenient_amount(payment.get(k) or 0, f"第{i + 1}期金额", draft, issues)
                   for i, k in enumerate(("pay1", "pay2", "pay3"))]
        for i, (row, amt) in enumerate(zip(inst_tbl.rows, amounts)):
            _imm_set_row(row, preset[i]["label"], f"于{preset[i]['event']}", amt, cur, rate)
        n = len(preset)
    elif mode == "one_time":
        desc = _ot_desc(payment, draft, issues) or PLACEHOLDER_DATE
        amounts = [total]
        _imm_set_row(inst_tbl.rows[0], "第一期", f"于{desc}一次性支付", total, cur, rate)
        for idx in range(len(inst_tbl.rows) - 1, 0, -1):
            W.delete_row(inst_tbl, idx)
        n = 1
    else:  # custom
        installments = payment.get("installments") or []
        if not installments:
            raise ValueError("自定义分期缺少期次数据")
        amounts = [_lenient_amount(x.get("amount") or 0, f"第{i + 1}期金额", draft, issues)
                   for i, x in enumerate(installments)]
        if any(a <= 0 for a in amounts):
            msg = "各期金额必须大于 0"
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
            amounts = [max(a, 0) for a in amounts]
        if sum(amounts) != total:
            msg = f"各期合计 {sum(amounts)} ≠ 总费用 {total}，请先在界面修正"
            if not draft:
                raise ValueError(msg)
            issues.append(msg)
        for i, ins in enumerate(installments):
            if not (W.is_valid_iso_date((ins.get("trigger_date") or "").strip())
                    or (ins.get("trigger") or "").strip()):
                msg = f"第{i + 1}期缺少付款时间／节点／条件（既无日期也无触发条件），不能生成"
                if not draft:
                    raise ValueError(msg)
                issues.append(msg)
        # 以模板首行为样式源克隆到 n 行，逐行改写散文句（首行即第一期，克隆补其余期数）
        while len(inst_tbl.rows) > 1:
            W.delete_row(inst_tbl, 1)
        for _ in range(len(installments) - 1):
            W.clone_row(inst_tbl, 0)
        for i, ins in enumerate(installments):
            when = payment_when(ins) or PLACEHOLDER_DATE
            _imm_set_row(inst_tbl.rows[i], _row_name(installments, i),
                         f"于{pay_when_phrase(when)}", amounts[i], cur, rate)
        n = len(installments)

    phase = fee_tbl.rows[0].cells[cfg.get("phase_cell", 3)]
    # 中文量词习惯：两期（不是「二期」）
    want = "一次性支付：" if mode == "one_time" else f"分{'两' if n == 2 else W.seq_cn(n)}期支付："
    if phase.text.strip() != want:
        W.set_cell_text(phase, want)
    # fee_rows 语义=费用表（总费用行表）行标签，供 checker 指纹比对；分期散文行由
    # check_rules 的移民分支直接从表格解析金额校验合计
    report["fee_rows"] = [W.row_label(r) for r in fee_tbl.rows]
    report["amounts"] = [total] + amounts
    report["section_paras"] = None
