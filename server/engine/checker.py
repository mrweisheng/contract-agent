# -*- coding: utf-8 -*-
"""
核对器：
1) 规则引擎（硬校验）：必填/残留空位/金额一致/分期合计/币种唯一/编号格式
2) 格式指纹：生成文件 vs 模板，逐段逐表比对样式与文字，非修改点必须零变化
"""
import re
from docx import Document

from engine import writer as W
from engine import car_extras
from engine.money import to_cn_upper, currency_label, currency_symbol
from engine.types_config import get_type

BLANK_RE = re.compile(r"_{2,}")
CUR_TOKENS = ["HK$", "¥", "$", "港币", "人民币"]


# ---------------- 文本提取 ----------------

def extract_contract_text(path_or_doc) -> str:
    if isinstance(path_or_doc, str):
        doc = Document(path_or_doc)
    else:
        doc = path_or_doc
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for i, t in enumerate(doc.tables):
        lines.append(f"〔表格{i + 1}〕")
        for row in t.rows:
            cells = []
            for c in row.cells:
                if c.text.strip() not in cells:
                    cells.append(c.text.strip())
            lines.append(" | ".join(cells))
    return "\n".join(lines)


# ---------------- 规则引擎 ----------------

def check_rules(type_key: str, form: dict, payment: dict, report: dict, doc) -> list:
    """返回硬错误列表（空列表=通过）。"""
    cfg = get_type(type_key)
    errors = []

    # 1. 必填字段
    for g in cfg["groups"]:
        for f in g["fields"]:
            if f.get("required") and f.get("auto"):
                continue
            if f.get("required") and not str(form.get(f["key"]) or "").strip():
                errors.append(f"必填项缺失：{f['label']}")
    for f in cfg["fee_fields"]:
        if f.get("required") and not form.get(f["key"]):
            errors.append(f"必填项缺失：{f['label']}")
    if cfg.get("port_required") and not form.get("port"):
        errors.append("必填项缺失：口岸")

    # 2. 残留空位（编号行/双方表/数据表/费用表/条款区）
    meta = doc.paragraphs[0].text
    if BLANK_RE.search(meta):
        errors.append("顶部编号/签署日期行仍有空位未填")

    fee_tbl = W.find_table(doc, cfg["fee_table_anchor"])
    for tbl_anchor in [cfg["party_table_anchor"]] + [kv["anchor"] for kv in cfg.get("kv_tables", [])]:
        t = W.find_table(doc, tbl_anchor)
        if t:
            for row in t.rows:
                joined = " ".join(c.text for c in row.cells)
                if BLANK_RE.search(joined):
                    errors.append(f"信息表存在未填空位：{row.cells[0].text.strip()[:20]}")
    if fee_tbl:
        for row in fee_tbl.rows:
            joined = " ".join(c.text for c in row.cells)
            if BLANK_RE.search(joined):
                errors.append(f"费用表存在未填空位：{row.cells[0].text.strip()[:20]}")
    if cfg.get("exchange_fee"):
        # 行名可能已随客户用词改写（上车费用），按锚点+别名依次找
        xf = cfg["exchange_fee"]
        t = None
        for name in [xf["anchor"], *xf.get("aliases", [])]:
            t = W.find_table(doc, name)
            if t is not None:
                break
        if t and BLANK_RE.search(" ".join(c.text for c in t.rows[0].cells)):
            errors.append("换车费用表存在未填空位")
    # 条款区残留空位
    sec_no, next_no = cfg["payment_section"]
    try:
        s, e = W.section_range(doc, sec_no, next_no)
        for p in doc.paragraphs[s:e]:
            if BLANK_RE.search(p.text):
                errors.append(f"付款条款区仍有空位未填：{p.text[:30]}…")
    except ValueError as ex:
        errors.append(str(ex))

    # 3. 金额一致性（从文档解析）
    cur = report["currency"]
    if fee_tbl:
        # 总额行
        row0 = fee_tbl.rows[0]
        num, cn = _parse_amount_cells(row0, cfg)
        if num is None:
            errors.append("费用表总金额行无法解析金额")
        else:
            if num != report["total"]:
                errors.append(f"总金额不一致：文档 {num} ≠ 表单 {report['total']}")
            if cn and cn != to_cn_upper(num):
                errors.append("总金额大写与数字不符")
        # 分期合计（不含总金额行；无金额的行如日期行自然跳过）
        amounts = []
        for row in fee_tbl.rows[1:]:
            n2, _ = _parse_amount_cells(row, cfg)
            if n2 is not None and n2 > 0:
                amounts.append(n2)
        if amounts and report.get("mode") in ("default", "custom"):
            if sum(amounts) != report["total"]:
                errors.append(f"分期合计 {sum(amounts)} ≠ 总费用 {report['total']}")

    # 4. 币种唯一。费用表数字写法带设计内第二币种换算后缀（「 / ¥X」或「 / HK$X」，
    #    固定汇率、方向与主币种相反），扫描前剥掉；只剥预期方向——方向写反仍算违规
    other_sym = currency_symbol("CNY" if cur == "HKD" else "HKD")
    conv_suffix_re = re.compile(r"\s*/\s*" + re.escape(other_sym) + r"[\d,]+")
    tokens = set()
    if fee_tbl:
        for row in fee_tbl.rows:
            for c in row.cells:
                text = conv_suffix_re.sub("", c.text)
                for tk in ("HK$", "港币"):
                    if tk in text:
                        tokens.add("HKD")
                for tk in ("¥", "人民币"):
                    if tk in text:
                        tokens.add("CNY")
                # 新办模板残留的孤立 $ 归为港币语义
                if re.search(r"(?<![A-Za-z￥])\$[\d,]", text):
                    tokens.add("HKD")
    if tokens and tokens != {cur}:
        errors.append(f"币种不唯一：文档中出现 {sorted(tokens)}，与所选 {cur} 不一致")

    # 5. 编号格式（自动=11位数字；手写放宽为8~14位数字）
    no = report.get("agreement_no", "")
    if not re.fullmatch(r"\d{8,14}", no):
        errors.append(f"合约编号格式异常：{no}")

    # 6. 卖车可选内容：质保时长/公里数可解析（builder 已拦，此处纵深防御）
    if type_key == "car" and car_extras.warranty_on(form):
        try:
            car_extras.warranty_parts(form)
        except ValueError as ex:
            errors.append(str(ex))

    return errors


def _parse_amount_cells(row, cfg):
    """从费用行解析（数字金额, 大写金额含「元整」）。"""
    texts = [c.text for c in row.cells]
    num = cn = None
    for t in texts:
        m = re.search(r"(?:HK\$|¥|\$)\s*([\d,]+)", t)
        if m and num is None:
            num = int(m.group(1).replace(",", ""))
        m2 = re.search(r"（大写）(.+?元整)", t)
        if m2 and cn is None:
            cn = m2.group(1)
    return num, cn


# ---------------- 格式指纹 ----------------

def _fmt_of(r):
    """单个 run 的格式签名（字体/字号/加粗/颜色）。"""
    f = r.font
    color = None
    try:
        color = str(f.color.rgb) if f.color and f.color.rgb else None
    except Exception:
        color = None
    return (f.name, str(f.size) if f.size else None, bool(f.bold), color)


def _run_sig(p):
    sigs = []
    for r in p.runs:
        sigs.append(_fmt_of(r))
    return (p.style.name, str(p.alignment), tuple(sigs))


def _body_run_sig(p):
    """段落正文 run（最长非空 run）的格式签名——文字应挂在它上面。
    指纹只比对 run 格式序列不看文字归属，文字若写进加粗前缀等特殊 run
    序列不变但视觉已跑偏，此签名专抓该类问题。"""
    runs = [r for r in p.runs if r.text]
    if not runs:
        return None
    return _fmt_of(max(runs, key=lambda r: len(r.text)))


def _fmt_multiset(row):
    """一行内所有非空 run 的格式计数（比对生成行是否出现模板没有的格式）。"""
    from collections import Counter
    c = Counter()
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:
                if r.text:
                    c[_fmt_of(r)] += 1
    return c


def _tbl_texts(t):
    out = []
    for row in t.rows:
        out.append(tuple(c.text for c in row.cells))
    return out


def _car_expected_tail(tp, te, form):
    """合并版卖车模板尾部期望段落序列（中文条款号）。
    条件块：四节尾「甲方陈述并保证」块（accident_free）/ 五节附赠服务（gift_*）/
    六节售后质保（warranty_*）——有内容保留并按推导文字，无内容整块不出现。
    固定节标题因删块重排，期望文字＝模板标题换前导中文序号。
    条目：("head", 期望文字, 模板下标) / ("same", 模板下标) / ("text", 期望文字, 模板下标)。"""
    gifts = car_extras.gift_lines(form)
    wp = car_extras.warranty_paragraph(form)
    accident = car_extras.accident_free_on(form)
    try:
        period, km = car_extras.warranty_custom_parts(form)
    except ValueError:
        period = km = None
    rewrite_passage = car_extras.gift_transfer_on(form)
    rewrite_ins = bool(car_extras.gift_insurance(form))

    def renum(t, n):
        return re.sub(r"^[一二三四五六七八九十]{1,3}", W.seq_cn(n), t, count=1)

    # 模板尾部标题序列 + 条件区范围
    heads = [(i, tp[i].text.strip()) for i in range(te, len(tp))
             if W.HEADING_RE.match(tp[i].text.strip())]
    def span(marker):
        for k, (i, t) in enumerate(heads):
            if marker in t:
                return i, (heads[k + 1][0] if k + 1 < len(heads) else len(tp))
        return None, None
    gi, gj = span("附赠服务")
    wi, wj = span("售后质保")
    bi = bj = None
    for i in range(te, len(tp)):
        if tp[i].text.strip() == "甲方陈述并保证：":
            bi = i
            for j in range(i + 1, len(tp)):
                if W.HEADING_RE.match(tp[j].text.strip()):
                    bj = j
                    break
            break

    exp = []
    n = 3          # te 处为四节标题，进入循环后递增
    for i in range(te, len(tp)):
        t = tp[i].text.strip()
        if W.HEADING_RE.match(t):
            if gi is not None and gi <= i < gj:
                if gifts:
                    n += 1
                    exp.append(("head", renum(t, n), i))
                    exp.extend(("text", g, i + 1) for g in gifts)
                continue
            if wi is not None and wi <= i < wj:
                if wp:
                    n += 1
                    exp.append(("head", renum(t, n), i))
                continue
            n += 1
            exp.append(("head", renum(t, n), i))
            continue
        if bi is not None and bi <= i < bj:
            if accident:
                exp.append(("same", i))
            continue
        if gi is not None and gi < i < gj:
            continue   # 占位行已由 gift 行代替
        if wi is not None and wi < i < wj:
            if not wp:
                continue   # 质保节整节删除（含节尾空段）
            if "保修范围" in t:
                txt = t
                if period:
                    txt = txt.replace("六个月", period)
                if km:
                    txt = txt.replace("三万公里", km)
                exp.append(("text", txt, i))
            else:
                exp.append(("same", i))
            continue
        if rewrite_passage and "车辆过户手续由甲方负责办理" in t:
            exp.append(("text", car_extras.passage_text(form), i))
            continue
        if rewrite_ins and "现有保险不随车转移" in t:
            exp.append(("text", car_extras.insurance_text(form), i))
            continue
        exp.append(("same", i))
    return exp


def check_fingerprint(template_path: str, out_path: str, type_key: str, report: dict,
                      form: dict = None) -> list:
    """比对模板与生成件：非修改点必须零变化；修改点样式签名必须与源一致。
    form（卖车）传入时按表单推导可选内容（附赠/质保）的期望落点与文字。"""
    cfg = get_type(type_key)
    issues = []
    tpl, out = Document(template_path), Document(out_path)

    # 页眉页脚、节设置
    if len(tpl.sections) != len(out.sections):
        issues.append("节数量与模板不一致")
    else:
        for st, so in zip(tpl.sections, out.sections):
            if st.header is not None and so.header is not None:
                ht = "\n".join(p.text for p in st.header.paragraphs)
                ho = "\n".join(p.text for p in so.header.paragraphs)
                if ht != ho:
                    issues.append("页眉内容与模板不一致")
            ft = "\n".join(p.text for p in st.footer.paragraphs)
            fo = "\n".join(p.text for p in so.footer.paragraphs)
            if ft != fo:
                issues.append("页脚内容与模板不一致")

    # 正文段落：条款区前的所有段落必须完全一致
    sec_no, next_no = cfg["payment_section"]
    try:
        ts, te = W.section_range(tpl, sec_no, next_no)
        os_, oe = W.section_range(out, sec_no, next_no)
    except ValueError as ex:
        issues.append(f"找不到条款区：{ex}")
        return issues

    tp, op = tpl.paragraphs, out.paragraphs
    for i in range(ts - 1):
        if i == 0:
            continue  # 首段为编号/日期行（设计内填充）
        if i >= len(op):
            issues.append(f"段落缺失（第{i + 1}段）")
            break
        if tp[i].text != op[i].text:
            issues.append(f"条款区之前段落被改动：第{i + 1}段「{tp[i].text[:20]}…」")
        elif _run_sig(tp[i]) != _run_sig(op[i]):
            issues.append(f"条款区之前段落格式变化：第{i + 1}段")
    # 条款区之后（尾部对齐；卖车开启附赠/质保时按表单推导的期望序列比对）
    tpl_tail = tp[te:]
    out_tail = op[oe:]
    expected = None
    if type_key == "car" and form is not None:
        try:
            expected = _car_expected_tail(tp, te, form)
        except ValueError as ex:
            issues.append(str(ex))
    if expected is not None:
        if len(out_tail) != len(expected):
            issues.append(f"条款区之后段落数量与期望不一致（{len(out_tail)} ≠ {len(expected)}）")
        else:
            for b, exp in zip(out_tail, expected):
                kind = exp[0]
                if kind == "same":
                    a = tp[exp[1]]
                    if a.text != b.text:
                        issues.append(f"条款区之后段落被改动：「{a.text[:20]}…」")
                        break
                    if _run_sig(a) != _run_sig(b):
                        issues.append("条款区之后段落格式变化")
                        break
                elif kind == "head":
                    _, txt, base = exp
                    if b.text != txt:
                        issues.append(f"条款标题与期望不符：「{b.text[:24]}…」")
                        break
                    if _run_sig(b) != _run_sig(tp[base]):
                        issues.append(f"条款标题样式与模板不一致：「{b.text[:24]}…」")
                        break
                else:  # text：条件节正文 / 过户费用句改写
                    _, txt, base = exp
                    if b.text != txt:
                        issues.append(f"可选条款文字与表单不符：「{b.text[:24]}…」")
                        break
                    if _run_sig(b) != _run_sig(tp[base]):
                        issues.append("可选条款段落样式与模板不一致")
                        break
    elif len(tpl_tail) != len(out_tail):
        issues.append("条款区之后段落数量与模板不一致")
    elif any(a.text != b.text for a, b in zip(tpl_tail, out_tail)):
        for a, b in zip(tpl_tail, out_tail):
            if a.text != b.text:
                issues.append(f"条款区之后段落被改动：「{a.text[:20]}…」")
                break
    else:
        for a, b in zip(tpl_tail, out_tail):
            if _run_sig(a) != _run_sig(b):
                issues.append("条款区之后段落格式变化")
                break
    # 条款区：段数=预期；改写块签名=模板首段，且文字必须落在正文 run 上；
    # 未改写（default）时逐段与模板同序比对
    if report.get("section_paras"):
        n_expect = report["section_paras"]
        if oe - os_ != n_expect:
            issues.append(f"付款条款段数 {oe - os_} ≠ 预期 {n_expect}")
        rw = report.get("payment_rewrite")
        if rw:
            a, b = rw
            seg = op[os_ + a:os_ + b]
            base_sig = _run_sig(tp[ts])
            for p in seg:
                if _run_sig(p) != base_sig:
                    issues.append("付款条款段落样式与模板不一致")
                    break
            body = _body_run_sig(tp[ts])
            if body:
                for p in seg:
                    if p.text.strip() and _body_run_sig(p) != body:
                        issues.append("付款条款正文格式与模板不一致（文字落到了特殊格式 run）")
                        break
        else:
            for p, q in zip(tp[ts:te], op[os_:oe]):
                if _run_sig(p) != _run_sig(q):
                    issues.append("付款条款段落样式与模板不一致")
                    break

    # 表格：除设计内会变的表（甲乙双方/数据表/费用表/换车费表）外，其余必须一致
    skip_anchors = [cfg["fee_table_anchor"], cfg["party_table_anchor"]]
    skip_anchors += [kv["anchor"] for kv in cfg.get("kv_tables", [])]
    if cfg.get("exchange_fee"):
        skip_anchors.append(cfg["exchange_fee"]["anchor"])

    def _tbl_alltext(t):
        return " ".join(c.text for r in t.rows for c in r.cells)

    def _norm_dates(s):
        """签署表日期归一化：模板空位与填充后的日期视为同一设计内变更。"""
        return re.sub(r"日期：[^年]{1,12}年[^月]{1,6}月[^日]{1,6}日", "日期：<日期>", s)

    for t_t in tpl.tables:
        t_all = _tbl_alltext(t_t)
        if "授权代表签署" in t_all:
            # 签署确认表：日期空位填充属设计内变更（归一化后比对），其余内容必须与模板一致。
            # 必须在 skip_anchors 之前判定——签署表首格含「甲方（卖方）」会被当事表锚点误跳过。
            t_o = next((c for c in out.tables if "授权代表签署" in _tbl_alltext(c)), None)
            if t_o is None:
                issues.append("表格丢失：签署确认块")
                continue
            a = tuple(tuple(_norm_dates(c.text) for c in r.cells) for r in t_t.rows)
            b = tuple(tuple(_norm_dates(c.text) for c in r.cells) for r in t_o.rows)
            if a != b:
                issues.append("表格内容被意外改动：签署确认块")
            continue
        if any(x in t_all for x in skip_anchors):
            continue
        t_first = t_t.rows[0].cells[0].text.strip()
        t_o = None
        for cand in out.tables:
            if t_first and t_first in _tbl_alltext(cand):
                t_o = cand
                break
        if t_o is None:
            issues.append(f"表格丢失：{t_first[:15]}")
            continue
        if _tbl_texts(t_t) != _tbl_texts(t_o):
            issues.append(f"表格内容被意外改动：{t_first[:15]}")

    # 费用表：custom 日期列模式下允许列数 +1；行样式签名=模板同标签行（无同标签时回退模板第1数据行）；标签符合预期
    tpl_fee = W.find_table(tpl, cfg["fee_table_anchor"])
    out_fee = W.find_table(out, cfg["fee_table_anchor"])
    if tpl_fee is None or out_fee is None:
        issues.append("找不到费用表")
        return issues
    extra_cols = 1 if report.get("date_column") else 0
    if len(tpl_fee.columns) + extra_cols != len(out_fee.columns):
        issues.append("费用表列数与模板不一致")

    def _row_sig(row, trim=0):
        cells = row.cells[: len(row.cells) - trim] if trim else row.cells
        return tuple(_run_sig(p) for c in cells for p in c.paragraphs)

    sig_by_label = {r.cells[0].text.strip(): _row_sig(r) for r in tpl_fee.rows}
    fallback_sig = _row_sig(tpl_fee.rows[1])
    for row in out_fee.rows:
        label = row.cells[0].text.strip()
        expect = sig_by_label.get(label, fallback_sig)
        if _row_sig(row, extra_cols) != expect:
            issues.append(f"费用表行样式与模板不一致：{label[:15]}")
            break
    expect_labels = report.get("fee_rows") or []
    actual_labels = [r.cells[0].text.strip() for r in out_fee.rows]
    if expect_labels and actual_labels != expect_labels:
        issues.append(f"费用表行标签与预期不符：{actual_labels}")
    # 生成行非空 run 的格式必须都在模板同标签行（无同标签回退第1数据行）出现过，
    # 否则说明写入文字挂到了模板中不存在的格式上（文字落点跑偏）
    ms_by_label = {r.cells[0].text.strip(): _fmt_multiset(r) for r in tpl_fee.rows}
    fb_ms = _fmt_multiset(tpl_fee.rows[1])
    for row in out_fee.rows:
        label = row.cells[0].text.strip()
        expect_ms = ms_by_label.get(label, fb_ms)
        got_ms = _fmt_multiset(row)
        if got_ms and got_ms - expect_ms:
            extra = sorted(got_ms - expect_ms)
            issues.append(f"费用表出现模板中没有的格式：{label[:15]} {extra}")
            break

    return issues
