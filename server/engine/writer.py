# -*- coding: utf-8 -*-
"""
docx 落盘执行器：所有修改都在原模板副本上做最小化编辑，
继承原有 run 格式（字体/字号/加粗/颜色），保证排版零变化。
"""
import copy
import re
from docx import Document
from docx.oxml.ns import qn

BLANK_RE = re.compile(r"_{2,}")


# ---------------- 段落级操作 ----------------

def para_text(p) -> str:
    return "".join(r.text for r in p.runs)


def replace_in_paragraph(p, old: str, new: str) -> bool:
    """跨 run 查找替换：新文本放入首个受影响 run（继承其格式）。"""
    full = "".join(r.text for r in p.runs)
    idx = full.find(old)
    if idx < 0:
        return False
    pos = 0
    first = True
    for r in p.runs:
        rs, re_ = pos, pos + len(r.text)
        pos = re_
        if re_ <= idx or rs >= idx + len(old):
            continue
        pre = r.text[: max(0, idx - rs)]
        post = r.text[max(0, min(len(r.text), idx + len(old) - rs)):]
        if first:
            r.text = pre + new + post
            first = False
        else:
            r.text = pre + post
    return True


def fill_blanks_in_paragraph(p, values) -> int:
    """按顺序把段落中连续下划线空位替换为 values（日期空「____年__月__日」即三空）。"""
    n = 0
    for v in values:
        full = "".join(r.text for r in p.runs)
        m = BLANK_RE.search(full)
        if not m:
            break
        _splice(p, m.start(), m.end(), str(v))
        n += 1
    return n


def _splice(p, start: int, end: int, new: str):
    pos = 0
    first = True
    for r in p.runs:
        rs, re_ = pos, pos + len(r.text)
        pos = re_
        if re_ <= start or rs >= end:
            continue
        pre = r.text[: max(0, start - rs)]
        post = r.text[max(0, min(len(r.text), end - rs)):]
        if first:
            r.text = pre + new + post
            first = False
        else:
            r.text = pre + post


def set_paragraph_text(p, text: str):
    """整段替换文字，保留格式：写入承载正文最长的 run，清空其余。
    不写首个非空 run——模板段落常是「加粗前缀：」+ 正文两个 run
    （如「第一期定金：」加粗深蓝、「于本合约签订当日支付…」常规），
    写首个 run 会让整句继承前缀的加粗/颜色。"""
    runs = [r for r in p.runs]
    if not runs:
        p.add_run(text)
        return
    target = None
    for r in runs:
        if r.text and (target is None or len(r.text) > len(target.text)):
            target = r
    if target is None:
        target = runs[0]
    for r in runs:
        if r is not target:
            r.text = ""
    target.text = text


def find_paragraph(doc, contains: str):
    for p in doc.paragraphs:
        if contains in p.text:
            return p
    return None


def _is_heading(p, no: str) -> bool:
    t = p.text.strip()
    return bool(re.match(rf"^{no}[\s　]", t)) or t.startswith(no)


def section_range(doc, sec_no: str, next_no: str):
    """返回（起始下标, 结束下标）之间的正文段落（不含标题）。"""
    start = end = None
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if start is None and _is_heading(p, sec_no):
            start = i + 1
            continue
        if start is not None and _is_heading(p, next_no):
            end = i
            break
    if start is None:
        raise ValueError(f"找不到条款 {sec_no} 标题")
    if end is None:
        end = len(paras)
    return start, end


def _is_payment_para(p, labels) -> bool:
    """段首 12 字内含款项名称（定金/第二期款/尾款…）→ 分期付款描述段。"""
    head = p.text.strip()[:12]
    return any(lb in head for lb in labels)


def rewrite_section(doc, sec_no: str, next_no: str, texts, labels=None):
    """改写条款区付款段：段落数量调整为 len(texts) 并逐段写入（克隆/删除保持格式）。
    labels（款项名称列表）给定时只改写「段首含款项名称」的连续块——
    新办模板 03 区还含费用范围/逾期付款/收款账户等无关条款，必须原样保留。
    返回（条款区总段数, 改写块在条款区内的起止下标）。"""
    sec = section_range(doc, sec_no, next_no)
    start, end = _shrink_to_payment_block(doc, sec, labels)
    paras = doc.paragraphs
    cur = paras[start:end]
    n_need = len(texts)
    n_have = len(cur)
    if n_need > n_have:
        last = cur[-1]._p
        for _ in range(n_need - n_have):
            new_p = copy.deepcopy(last)
            last.addnext(new_p)
            last = new_p
    elif n_need < n_have:
        for p in cur[n_need:]:
            p._p.getparent().remove(p._p)
    # 重新抓取（索引已变）
    sec2 = section_range(doc, sec_no, next_no)
    start2, end2 = _shrink_to_payment_block(doc, sec2, labels)
    for p, t in zip(doc.paragraphs[start2:end2], texts):
        set_paragraph_text(p, t)
    total = sec2[1] - sec2[0]
    return total, start2 - sec2[0], end2 - sec2[0]


def _shrink_to_payment_block(doc, sec, labels):
    """把条款区 (start, end) 收缩到分期付款段块；无 labels 或无命中时原样返回。"""
    start, end = sec
    if not labels:
        return start, end
    hits = [i for i in range(start, end) if _is_payment_para(doc.paragraphs[i], labels)]
    if not hits:
        return start, end
    return hits[0], hits[-1] + 1


# ---------------- 表格级操作 ----------------

def iter_tables(doc):
    return list(doc.tables)


def find_table(doc, anchor: str):
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if anchor in c.text:
                    return t
    return None


def cell_text(cell) -> str:
    return cell.text.strip()


def set_cell_text(cell, text: str):
    """单元格写文字（首个段落），保留原格式。"""
    p = cell.paragraphs[0]
    set_paragraph_text(p, text)
    # 单元格可能有多个段落，多余段落清空（模板单元格均为单段）
    for extra in cell.paragraphs[1:]:
        set_paragraph_text(extra, "")


def clone_row(table, src_idx: int):
    """克隆第 src_idx 行（含格式）追加到表尾，返回新行。"""
    src_tr = table.rows[src_idx]._tr
    new_tr = copy.deepcopy(src_tr)
    # 去掉克隆行的 paraId 重复
    for el in new_tr.iter():
        for attr in list(el.attrib):
            if attr.endswith("}paraId") or attr.endswith("}textId"):
                del el.attrib[attr]
    table._tbl.append(new_tr)
    return table.rows[-1]


def payment_row_name(labels, i: int) -> str:
    """分期行名：客户对这笔款项的称呼（定金/尾款…）优先，
    空／超 6 字／与其他期重名时退回「第X期」兜底。
    解析回填（routes）与落盘（builder）共用，保证页面所见即合同所得。"""
    lab = str(labels[i] or "").strip()
    if lab and len(lab) <= 6 and labels.count(lab) == 1:
        return lab
    return f"第{seq_cn(i + 1)}期"


WIDE_COL_TWIPS = 2000  # 宽于此的列才参与让位压缩（窄标签列不动）


def append_column(table, width: int):
    """表尾追加一列（固定布局表格）：克隆每行末列单元格以继承边框/字体，
    列宽从原宽列等比腾出（总宽不变，不撑破版心）。返回各行新列的 cell。"""
    grid = table._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    ws = [int(c.get(qn("w:w"))) for c in cols]
    # 只压宽列，标签窄列原样保留
    wide = [i for i, w in enumerate(ws) if w > WIDE_COL_TWIPS]
    if not wide or sum(ws[i] for i in wide) < width:
        raise ValueError("费用表没有可腾出的宽列，无法追加付款日期列")
    freed = {i: ws[i] * width // sum(ws[j] for j in wide) for i in wide}
    drift = width - sum(freed.values())
    freed[max(wide, key=lambda i: freed[i])] += drift
    for i in wide:
        cols[i].set(qn("w:w"), str(ws[i] - freed[i]))
    new_gc = copy.deepcopy(cols[-1])
    new_gc.set(qn("w:w"), str(width))
    grid.append(new_gc)
    for tr in table._tbl.tr_lst:
        tcs = tr.findall(qn("w:tc"))
        new_tc = copy.deepcopy(tcs[-1])
        for el in new_tc.iter():
            for attr in list(el.attrib):
                if attr.endswith("}paraId") or attr.endswith("}textId"):
                    del el.attrib[attr]
        tcPr = new_tc.find(qn("w:tcPr"))
        tcW = tcPr.find(qn("w:tcW")) if tcPr is not None else None
        if tcW is not None:
            tcW.set(qn("w:w"), str(width))
        tr.append(new_tc)
    return [row.cells[-1] for row in table.rows]


def delete_row(table, idx: int):
    tr = table.rows[idx]._tr
    tr.getparent().remove(tr)


def row_label(row) -> str:
    return row.cells[0].text.strip()


# ---------------- 日期工具 ----------------

def is_valid_iso_date(s: str) -> bool:
    """YYYY-M-D 且为真实日历日期（LLM 可能给出 2026-02-30 之类幻觉日期）。"""
    import datetime
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", str(s or "").strip())
    if not m:
        return False
    try:
        datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        return True
    except ValueError:
        return False


def iso_date_parts(iso: str):
    y, m, d = iso.split("-")
    return y, str(int(m)), str(int(d))


def date_cn(iso: str) -> str:
    y, m, d = iso_date_parts(iso)
    return f"{y}年{m}月{d}日"


def today_cn():
    import datetime
    t = datetime.date.today()
    return f"{t.year}年{t.month}月{t.day}日"


_CN_SEQ = "零一二三四五六七八九"


def seq_cn(i: int) -> str:
    if 1 <= i <= 10:
        return "一二三四五六七八九十"[i - 1]
    if i < 20:
        return "十" + (_CN_SEQ[i % 10] if i % 10 else "")
    if i < 100:
        t, u = divmod(i, 10)
        return _CN_SEQ[t] + "十" + (_CN_SEQ[u] if u else "")
    return str(i)
