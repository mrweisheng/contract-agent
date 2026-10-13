from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/chenqiqi/Documents/ChatGPT/cqtzVntTAhsv/outputs/settlement_template/客户收款结算单模板.docx"
LOGO = "/Users/chenqiqi/Documents/ChatGPT/cqtzVntTAhsv/assets/ming-ge-cross-border-logo-v1.png"

NAVY = "0B2E59"
BLUE = "1769AA"
PALE = "EDF4FA"
GOLD = "C99A3D"
GRAY = "687386"
BORDER = "D9E1EA"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(cell, color=BORDER, size="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)

def margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, color="1D2733", size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "STHeiti"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.62)
sec.right_margin = Inches(0.62)

styles = doc.styles
styles["Normal"].font.name = "STHeiti"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "STHeiti"
styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)

# Header
hdr = doc.add_table(rows=1, cols=2)
hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr.autofit = False
hdr.columns[0].width = Inches(3.5)
hdr.columns[1].width = Inches(3.65)
hdr.cell(0,0).width = Inches(3.5)
hdr.cell(0,1).width = Inches(3.65)
for c in hdr.rows[0].cells:
    borders(c, "FFFFFF", "0")
    margins(c, 0, 0, 0, 0)
p = hdr.cell(0,0).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run().add_picture(LOGO, width=Inches(2.2))
p = hdr.cell(0,1).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("SETTLEMENT STATEMENT\n")
r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(GRAY)
r2 = p.add_run("客户收款结算单")
r2.bold = True; r2.font.size = Pt(22); r2.font.color.rgb = RGBColor.from_string(NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(8)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("结算单编号  [自动生成]    出具日期  2026年9月21日")
r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(GRAY)

# Meta
meta = doc.add_table(rows=2, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
widths = [1.05, 2.55, 1.05, 2.55]
for row in meta.rows:
    for i, c in enumerate(row.cells):
        c.width = Inches(widths[i]); borders(c); margins(c, 105, 125, 105, 125)
labels = [("客户名称", "[客户名称]", "业务项目", "[项目或业务名称]"),
          ("结算类型", "[尾款／杂费／补款／阶段款]", "结算币种", "[人民币 CNY／港币 HKD]")]
for ri, vals in enumerate(labels):
    for ci, val in enumerate(vals):
        is_label = ci in (0,2)
        set_cell_text(meta.cell(ri,ci), val, bold=is_label, color="FFFFFF" if is_label else "1D2733", size=9.5)
        if is_label: shade(meta.cell(ri,ci), NAVY)

doc.add_paragraph().paragraph_format.space_after = Pt(1)

# Detail table
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [0.42, 2.45, 0.72, 1.12, 1.12, 1.32]
headers = ["序号", "款项说明", "币种", "应收金额", "已收或扣减", "本次应收"]
for i, w in enumerate(widths):
    table.columns[i].width = Inches(w)
    table.cell(0,i).width = Inches(w)
for i,h in enumerate(headers):
    borders(table.cell(0,i)); shade(table.cell(0,i), NAVY)
    set_cell_text(table.cell(0,i), h, True, "FFFFFF", 9, WD_ALIGN_PARAGRAPH.CENTER)

rows = [
    ("1", "业务合同金额或阶段款", "CNY", "¥ 0.00", "¥ 0.00", "¥ 0.00"),
    ("2", "订金或前期已付款", "CNY", "¥ 0.00", "− ¥ 0.00", "¥ 0.00"),
    ("3", "物流或仓储费用", "CNY", "¥ 0.00", "¥ 0.00", "¥ 0.00"),
    ("4", "保险 验车 校正本等杂费", "HKD", "HK$ 0.00", "HK$ 0.00", "HK$ 0.00"),
    ("5", "其他增加项", "CNY", "¥ 0.00", "¥ 0.00", "¥ 0.00"),
    ("6", "优惠 退款 抵扣或其他减少项", "CNY", "¥ 0.00", "− ¥ 0.00", "¥ 0.00"),
]
for ri, vals in enumerate(rows, 1):
    cells = table.add_row().cells
    for ci, val in enumerate(vals):
        cells[ci].width = Inches(widths[ci]); borders(cells[ci])
        if ri % 2 == 0: shade(cells[ci], "F7FAFD")
        align = WD_ALIGN_PARAGRAPH.LEFT if ci == 1 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(cells[ci], val, False, "1D2733", 9.2, align)

# Totals
for label, value, fill, color, size in [
    ("人民币本次应收合计", "¥ 0.00", PALE, NAVY, 11),
    ("港币本次应收合计", "HK$ 0.00", PALE, NAVY, 11),
    ("本次付款总额", "[按客户实际付款币种列示]", NAVY, "FFFFFF", 13),
]:
    cells = table.add_row().cells
    merged = cells[0].merge(cells[4])
    set_cell_text(merged, label, True, color, size, WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell_text(cells[5], value, True, color, size, WD_ALIGN_PARAGRAPH.CENTER)
    shade(merged, fill); shade(cells[5], fill)
    borders(merged); borders(cells[5])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("金额说明")
r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string(NAVY)
p.add_run("  正数表示本次新增应收，负数表示已付款、退款、优惠或其他扣减。人民币与港币分开核算，不混合折算；如需折算，将单独注明汇率和折算日期。")

# Payment section
pay = doc.add_table(rows=3, cols=4)
pay.alignment = WD_TABLE_ALIGNMENT.CENTER
pay.autofit = False
pay_widths = [1.05, 2.55, 1.05, 2.55]
pay_rows = [
    ("收款账户", "[账户名称]", "银行名称", "[开户银行]"),
    ("收款账号", "[银行账号]", "付款期限", "[YYYY年MM月DD日前]"),
    ("付款备注", "[客户名 项目名 结算单编号]", "联系人", "[联系人及联系方式]"),
]
for ri, vals in enumerate(pay_rows):
    for ci, val in enumerate(vals):
        c=pay.cell(ri,ci); c.width=Inches(pay_widths[ci]); borders(c)
        is_label=ci in (0,2)
        set_cell_text(c,val,is_label,"FFFFFF" if is_label else "1D2733",9.2)
        if is_label: shade(c,BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(9)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("请核对以上结算内容。若无异议，请按付款期限完成支付；付款后请将凭证发送给对接人员，以便及时核销入账。")
r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string(NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("明哥中港牌  MING GE CROSS BORDER  ·  商务结算文件")
r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(GRAY)

doc.save(OUT)
print(OUT)
